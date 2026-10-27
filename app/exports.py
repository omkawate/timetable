"""Horizontal timetable exports; scheduling and saved entries are not modified."""
import io
from pathlib import Path
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from html import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from .domain import COLLEGE
from .presentation import layouts,calendar_rows

FONT_DIR=Path(__file__).resolve().parent.parent/'static'/'fonts'
pdfmetrics.registerFont(TTFont('TimetableSerif-Bold',str(FONT_DIR/'DejaVuSerif-Bold.ttf')))
pdfmetrics.registerFont(TTFont('TimetableSans',str(FONT_DIR/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('TimetableSans-Bold',str(FONT_DIR/'DejaVuSans-Bold.ttf')))


def workflow_label(generation):
    w=generation.get('workflow',{'state':'draft','locked':False})
    return w['state'].upper()+(' · LOCKED' if w['locked'] else '')


def matrix(layout):
    data=[['']*(len(layout['columns'])+1) for _ in range(layout['row_count'])]
    data[0]=['DAYS']+[('RECESS' if 'break' in c else 'P'+str(c['slot']+1))+'\n'+c['time'] for c in layout['columns']]
    for cell in layout['cells']:data[cell['row']][cell['col']]=cell['text']
    return data


def views(generation,view='group',target=None,batch=None):
    return [(x['title'],matrix(x),x['hours']) for x in layouts(generation,view,target,batch)]


def excel(generation,view='group',target=None,batch=None):
    wb=Workbook();wb.remove(wb.active)
    for index,layout in enumerate(layouts(generation,view,target,batch),1):
        sheet=wb.create_sheet(f'{index} '+{'group':'Division','teacher':'Professor','classroom':'Classroom','lab':'Lab'}[view]);n=len(layout['columns'])+1
        border=Border(**{side:Side(style='thin',color='000000') for side in ['top','bottom','left','right']})
        def put(row,col,text,rs=1,cs=1,bold=False,size=11):
            for r in range(row,row+rs):
                for c in range(col,col+cs):sheet.cell(r,c).border=border
            if rs>1 or cs>1:sheet.merge_cells(start_row=row,start_column=col,end_row=row+rs-1,end_column=col+cs-1)
            cell=sheet.cell(row,col,text);cell.data_type='s';cell.font=Font(name='Times New Roman' if bold else 'Arial',size=size,bold=bold,color='000000')
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            return cell
        for row,text in enumerate(layout['headers'],1):
            put(row,1,text,cs=n,bold=True,size=16 if row==1 else 12);sheet.row_dimensions[row].height=24 if row==1 else 20
        bounds=[1,1+n//3,1+2*n//3,n+1]
        if len(layout['details'])==3:
            for i,text in enumerate(layout['details']):put(5,bounds[i],text,cs=bounds[i+1]-bounds[i],bold=True)
        else:put(5,1,layout['details'][0],cs=n,bold=True)
        sheet.row_dimensions[5].height=22
        put(6,1,'DAYS',rs=2,bold=True)
        for i,col in enumerate(layout['columns'],2):
            put(6,i,str(i-1),bold=True);put(7,i,col['time'],bold=True,size=10)
        sheet.row_dimensions[6].height=20;sheet.row_dimensions[7].height=20
        for r in range(8,layout['row_count']+7):sheet.row_dimensions[r].height=25
        for c in layout['cells']:
            text='\n'.join(c['text']) if c['kind']=='break' else c['text']
            put(c['row']+7,c['col']+1,text,c['rowspan'],c['colspan'],bold=c['kind'] in ('day','break'))
            if c['rowspan']==1 and c['text']:sheet.row_dimensions[c['row']+7].height=max(sheet.row_dimensions[c['row']+7].height,14*(c['text'].count('\n')+1)+8)
        start=layout['row_count']+7
        bounds=[1,1+n//4,1+n//2,1+3*n//4,n+1]
        # Match the four adjacent legend panels of the reference.
        legend_height=max(len(l['items']) for l in layout['legends'])+1
        for i,l in enumerate(layout['legends']):
            for offset in range(legend_height):
                text=l['title'] if offset==0 else l['items'][offset-1] if offset<=len(l['items']) else ''
                cell=put(start+offset,bounds[i],text,cs=bounds[i+1]-bounds[i],bold=offset==0,size=10)
                cell.alignment=Alignment(horizontal='center' if offset==0 else 'left',vertical='center',wrap_text=True)
                sheet.row_dimensions[start+offset].height=30
        end=start+legend_height
        bounds=[1,1+n//3,1+2*n//3,n+1]
        for i,text in enumerate(layout['signatures']):put(end,bounds[i],'\n\n'+text,cs=bounds[i+1]-bounds[i],bold=True)
        sheet.row_dimensions[end].height=90
        sheet.column_dimensions['A'].width=9
        for i,col in enumerate(layout['columns'],2):sheet.column_dimensions[get_column_letter(i)].width=10 if 'break' in col else 18
        sheet.freeze_panes='B8';sheet.sheet_view.showGridLines=False
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=1
        sheet.sheet_properties.pageSetUpPr.fitToPage=True;sheet.print_options.horizontalCentered=True
        sheet.print_area=f'A1:{get_column_letter(n)}{end}'
    out=io.BytesIO();wb.save(out);return out.getvalue()


def pdf(generation,view='group',target=None,batch=None):
    # Draw each form on one page; scaling applies only to unusually dense views.
    from reportlab.pdfgen import canvas
    out=io.BytesIO();page_size=landscape(A3);width=page_size[0]-40;height=page_size[1]-40
    canvas_=canvas.Canvas(out,pagesize=page_size);canvas_.setTitle('College timetable')
    heading=ParagraphStyle('heading',fontName='TimetableSerif-Bold',fontSize=15,leading=17,alignment=TA_CENTER)
    normal=ParagraphStyle('normal',fontName='TimetableSans',fontSize=10.5,leading=13,alignment=TA_CENTER)
    bold=ParagraphStyle('bold',parent=normal,fontName='TimetableSerif-Bold')
    legendstyle=ParagraphStyle('legend',fontName='TimetableSans',fontSize=9,leading=11)
    para=lambda text,style:Paragraph(escape(text).replace('\n','<br/>'),style)
    base=[('GRID',(0,0),(-1,-1),.6,colors.black),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3),('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3)]
    for layout in layouts(generation,view,target,batch):
        blocks=[]
        headers=Table([[para(t,heading if i==0 else bold)] for i,t in enumerate(layout['headers'])],colWidths=[width]);headers.setStyle(TableStyle(base));blocks.append(headers)
        details=Table([[para(t,bold) for t in layout['details']]],colWidths=[width/len(layout['details'])]*len(layout['details']));details.setStyle(TableStyle(base));blocks.append(details)
        weights=[.65]+[.72 if 'break' in c else 1 for c in layout['columns']];widths=[width*w/sum(weights) for w in weights]
        data=matrix(layout)
        data[0]=['DAYS']+[str(i+1) for i in range(len(layout['columns']))]
        data.insert(1,['']+[c['time'] for c in layout['columns']])
        grid=[[para(text,bold if r<2 else normal) for text in row] for r,row in enumerate(data)]
        commands=base+[('SPAN',(0,0),(0,1))]
        for c in layout['cells']:
            r=c['row']+1
            if c['kind']=='break':grid[r][c['col']]=para('\n'.join(c['text']),bold)
            if c['rowspan']>1 or c['colspan']>1:commands.append(('SPAN',(c['col'],r),(c['col']+c['colspan']-1,r+c['rowspan']-1)))
            if c['kind']=='day':grid[r][c['col']]=para(c['text'],bold)
        table=Table(grid,colWidths=widths,minRowHeights=[20,20]+[22]*(len(grid)-2));table.setStyle(TableStyle(commands));blocks.append(table)
        legends=[]
        for l in layout['legends']:legends.append([para(l['title'],bold)]+[para(x,legendstyle) for x in l['items']])
        legendrows=[[legends[c][r] if r<len(legends[c]) else '' for c in range(4)] for r in range(max(map(len,legends)))]
        legend=Table(legendrows,colWidths=[width/4]*4)
        legend.setStyle(TableStyle([('BOX',(0,0),(-1,-1),.6,colors.black),('LINEAFTER',(0,0),(2,-1),.6,colors.black),('LINEBELOW',(0,0),(-1,0),.6,colors.black),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),1),('BOTTOMPADDING',(0,0),(-1,-1),1)]));blocks.append(legend)
        signatures=Table([[para(t,bold) for t in layout['signatures']]],colWidths=[width/3]*3,rowHeights=[85]);signatures.setStyle(TableStyle(base+[('VALIGN',(0,0),(-1,-1),'BOTTOM'),('BOTTOMPADDING',(0,0),(-1,-1),9)]));blocks.append(signatures)
        heights=[b.wrap(width,height)[1] for b in blocks];scale=min(1,height/sum(heights))
        canvas_.saveState();canvas_.translate((page_size[0]-width*scale)/2,page_size[1]-20);canvas_.scale(scale,scale)
        y=0
        for block,h in zip(blocks,heights):y-=h;block.drawOn(canvas_,0,y)
        canvas_.restoreState();canvas_.showPage()
    canvas_.save();return out.getvalue()
